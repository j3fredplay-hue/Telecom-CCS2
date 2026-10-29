#!/usr/bin/env python3
"""Cruza as ONUs de uma OLT com os relatórios do IXC e gera os Excel de limpeza e de ativos.

Uso:
    python3 olt_ixc.py dados/olts/OLT1-SHOPMILLENNIUM.txt \
        --logins dados/ixc/Relatorio_com_Login.xlsx \
        --contratos dados/ixc/Relatorio_de_contrato_do_IXC.xlsx \
        --saida saida/

Regras:
  - O número após "CTR-" é procurado primeiro como login PPPoE e depois como ID de contrato.
  - Status de acesso Ativo, Bloqueio Automático/Manual e Financeiro em atraso contam como ATIVO.
  - Rede neutra: contrato com plano de venda 67 (Circuito Dedicado) ou número repetido em várias ONUs.
  - Parceiro Upwire (CTR-201939-...): o status vem do relatório do SGP dele (--sgp). Primeiro pelo
    número do contrato SGP (CTR-201939-11202), senão pelo nome do cliente (CTR-201939-MONICA).
  - Senhas e demais dados sensíveis do IXC nunca são copiados para a saída.
"""
import argparse
import collections
import functools
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PLANO_REDE_NEUTRA = 67
PARCEIRO_SGP = "201939"  # nº do sistema antigo da Upwire; no IXC é o contrato 388 (MARCELO SILVA DO CARMO)
NOME_PARCEIRO_SGP = "Upwire"
# Números do sistema antigo da CCS Brasil que não existem como login no IXC → ID do contrato no IXC.
NUMEROS_ANTIGOS = {
    "2023716": 538,  # ID HOUSE TECNOLOGIA E CERTIFICACAO LTDA
    "2025928": 364,  # VISION TELECOM SERVIÇO DE COMUNICAÇÃO MULTIMIDIA LTDA
}

# Contratos de provedores parceiros (rede neutra) — IXC, relatório "Porta - Rede Neutra" de 28/10/2026.
# Vale junto com o plano 67; inclui parceiros já desativados (as ONUs deles seguem o status do contrato).
PARCEIROS_EXTRA = {
    82, 119, 130, 131, 184, 233, 247, 296, 298, 299, 306, 317, 324, 325, 326, 336, 339, 351, 355, 356,
    357, 364, 375, 377, 382, 384, 389, 397, 400, 404, 408, 419, 455, 460, 476, 490, 491, 494, 510, 512,
    516, 523, 538, 541, 548, 563, 595, 599,
}
# Descrições só com nome → ID do contrato no IXC (confirmadas pelo usuário). Compara sem espaços/pontuação.
NOMES_CONHECIDOS = [
    ("PLANODESAUDEWINK", 247),  # cliente da WINK (rede neutra)
    ("SANTOSEVIDAL", 2),        # LABORATORIO SANTOS E VIDAL
    ("IDHOUSE", 538),           # ID HOUSE TECNOLOGIA E CERTIFICACAO
    ("STAFF", 64),              # STAFF CONSTRUCOES
    ("BOBS", 102),              # PN REDE SAUDAVEL (Bob's)
    ("CPDE", 189),              # CENTRO DE PESQUISA E DIAGNOSTICOS ESPECIALIZADOS
    ("NERDSFIBRA", 512),        # NERDS TELECOM (rede neutra)
    ("FECHACOM", 253),          # FECHACOM COM DE FECHADURAS
]

# Descrições sem contrato identificado que o usuário mandou remover, mesmo online.
REMOVER_SEM_CONTRATO = ["TINTAO"]

# Decisões pontuais do usuário, por (serial, descrição) da ONU: (classificação, motivo).
DECISOES_MANUAIS = {
    ("TPLGCE87505A", ""): ("Desativado", "ONU sem descrição — remover (confirmado pelo usuário)"),
    ("TPLGCE8750C1", ""): ("Desativado", "ONU sem descrição — remover (confirmado pelo usuário)"),
    ("ITBSE8D0ACDF", "CTR-TESTE"): ("Desativado", "ONU de teste — remover (confirmado pelo usuário)"),
    ("VSOL00CD28FB", "CTR-201939-REBECA"): ("Desativado", "CTR-201939-REBECA: ONU offline e PPPoE online por outro caminho — não usa esta OLT"),
}

SGP_ATIVO = {"ativo", "suspenso"}
SGP_DESATIVADO = {"cancelado"}
SGP_PRE = {"inativo"}  # contrato criado e nunca ativado
STATUS_ATIVO = {"ativo", "bloqueio automático", "bloqueio manual", "financeiro em atraso"}
STATUS_DESATIVADO = {"desativado"}
STATUS_PRE = {"aguardando assinatura"}

FONTE = "Arial"


# ---------------------------------------------------------------- leitura das OLTs

def ler_datacom(texto):
    """Saída de `show interface gpon 1/1/ onu` do DmOS."""
    padrao = re.compile(r"^(\d+/\d+/\d+)\s+(\d+)\s+(\S+)\s+(Up|Down)\s+\S+\s*(.*)$", re.I)
    onus = []
    for linha in texto.splitlines():
        m = padrao.match(linha.strip())
        if m:
            pon, onu_id, serial, estado, nome = m.groups()
            onus.append({"pon": pon, "onu_id": int(onu_id), "serial": serial,
                         "estado": estado.capitalize(), "descricao": nome.strip()})
    return onus


def ler_parks(texto):
    """Saída de `show interface gpon1/1 onu` (perfil) e/ou `... onu status` (estado) da Parks, no mesmo
    arquivo. Status ACTIVE = Up, INVALID = Down; sem o status, o estado fica desconhecido."""
    onus, pon, atual, por_id = [], None, None, {}
    for linha in texto.splitlines():
        m = re.match(r"^\s*Interface gpon(\d+/\d+)", linha)
        if m:
            pon = m.group(1)
            continue
        m = re.match(r"^\s{2,}(\d+)-(.*?)(?:\s*\(([0-9a-z]{12})\))?:\s*$", linha, re.I)
        if m and pon:
            nome, serial = m.group(2).strip(), m.group(3)
            if not serial and re.fullmatch(r"[0-9a-z]{12}", nome, re.I):
                nome, serial = "", nome  # ONU sem descrição: o nome é o próprio serial
            chave_onu = (pon, int(m.group(1)))
            atual = por_id.get(chave_onu)
            if not atual:
                atual = por_id[chave_onu] = {"pon": pon, "onu_id": int(m.group(1)), "serial": (serial or "").upper(),
                                             "estado": "Desconhecido", "descricao": nome}
                onus.append(atual)
            continue
        m = re.match(r"^\s*Status\s*:\s*(\S+)", linha)
        if m and atual:
            atual["estado"] = "Up" if m.group(1).upper() == "ACTIVE" else "Down"
            continue
        m = re.match(r"^\s*Power Level\s*:\s*(-?[\d.]+)dBm", linha)
        if m and atual:
            atual["sinal"] = float(m.group(1))
            continue
        m = re.match(r"^\s*Flow profile:\s*(\S+)", linha)
        if m and atual:
            atual["perfil"] = m.group(1)
    return onus


LEITORES = {"datacom": ler_datacom, "parks": ler_parks}


def detectar_fabricante(texto):
    if re.search(r"^\s*Interface gpon\d+/\d+", texto, re.M):
        return "parks"
    if "Software Download State" in texto or re.search(r"^\d+/\d+/\d+\s+\d+\s+", texto, re.M):
        return "datacom"
    raise SystemExit("Formato de OLT não reconhecido. Informe --fabricante.")


# ---------------------------------------------------------------- leitura do IXC

def ler_planilha(caminho):
    ws = openpyxl.load_workbook(caminho, read_only=True).active
    linhas = ws.iter_rows(values_only=True)
    cab = next(linhas)
    return [dict(zip(cab, r)) for r in linhas if any(v is not None for v in r)]


def chave(valor):
    return str(valor).strip().lower() if valor is not None else ""


class IXC:
    def __init__(self, arq_logins, arq_contratos):
        self.contratos = {int(r["ID"]): r for r in ler_planilha(arq_contratos)}
        self.logins = collections.defaultdict(list)
        for r in ler_planilha(arq_logins):
            self.logins[chave(r["Login"])].append(r)

    def contrato(self, cid):
        return self.contratos.get(int(cid)) if cid is not None else None


class SGP:
    """Relatório 'Contrato por Status' do SGP do parceiro (cabeçalho em 'Cliente ID')."""

    def __init__(self, caminho):
        ws = openpyxl.load_workbook(caminho, read_only=True).active
        linhas = list(ws.iter_rows(values_only=True))
        i = next(n for n, r in enumerate(linhas) if r and r[0] == "Cliente ID")
        cab = linhas[i]
        self.contratos = {}
        for r in linhas[i + 1:]:
            d = dict(zip(cab, r))
            if isinstance(d.get("Contrato ID"), int):
                d["Status"] = str(d["Status"] or "").strip()
                self.contratos[d["Contrato ID"]] = d
        self.por_cliente = collections.defaultdict(list)
        self.logins = collections.defaultdict(list)  # login aparece na observação da última mudança de status
        for c in self.contratos.values():
            self.por_cliente[c["Cliente ID"]].append(c)
            m = re.search(r"Login:\s*([^\s,]+)", str(c.get("Observação") or ""))
            if m:
                self.logins[chave(m.group(1)).split("@")[0]].append(c)

    def por_nome(self, texto):
        """Uma palavra: primeiro nome, ou primeiro nome + outro nome colados (DOUGLASVARELA).
        Várias palavras: todas precisam estar no nome."""
        palavras = [compacto(p) for p in re.split(r"[-_. ]+", texto) if len(compacto(p)) >= 3]
        if not palavras:
            return []
        achados = []
        for c in self.contratos.values():
            nome = [compacto(p) for p in str(c["Cliente"]).split()]
            if (len(palavras) == 1 and nome and (nome[0] == palavras[0] or
                                                 any(nome[0] + w == palavras[0] for w in nome[1:]))) or \
               (len(palavras) > 1 and all(p in nome for p in palavras)):
                achados.append(c)
        return achados


def nome_bate(desc, nome):
    """'joedizalima' bate com 'JOEDIZA SAMPAIO DE LIMA': começa pelo primeiro nome e o resto é formado
    por outros nomes (ou suas iniciais), na ordem, pulando os que faltarem."""
    w = [compacto(x) for x in str(nome).split() if compacto(x)]
    if not w or not desc.startswith(w[0]):
        return False
    if desc == w[0]:
        return len(desc) >= 8  # só o primeiro nome: aceita apenas nomes longos (ex.: empresas)

    @functools.lru_cache(None)
    def resto(pos, i):
        if pos == len(desc):
            return True
        return any(desc.startswith(t, pos) and resto(pos + len(t), j + 1)
                   for j in range(i, len(w)) for t in (w[j], w[j][0]))
    return resto(len(w[0]), 1)


def identificar_por_nome(o, ixc, sgp):
    """Descrição sem número (padrão Nova Solutec: login ou nome do cliente). Procura no IXC e no SGP Upwire:
    primeiro como login, depois pelo nome. Devolve True se resolveu a classificação."""
    texto = re.sub(r"^\s*C-?(?:TR|RT)\b", "", o["descricao"], flags=re.I)
    partes = [p for p in re.split(r"[-_ ]+", texto) if p]
    variantes = [compacto("".join(partes[i:])) for i in range(len(partes))]  # 'RN-CLICKSUFT-MARCIABULCAO'
    variantes = [v for v in dict.fromkeys(variantes) if len(v) >= 5]
    if not variantes:
        return False

    achados, por = [], "Login"
    for v in variantes:
        achados += [("IXC", r["ID contrato"]) for k, rs in ixc.logins.items() if k.split("@")[0] == v.lower() for r in rs]
        if sgp:
            achados += [("SGP", c["Contrato ID"]) for c in sgp.logins.get(v.lower(), [])]
        if achados:
            break
    if not achados:
        por = "Nome"
        for v in variantes:
            achados += [("IXC", cid) for cid, c in ixc.contratos.items() if nome_bate(v, c["Cliente"])]
            if sgp:
                achados += [("SGP", cid) for cid, c in sgp.contratos.items() if nome_bate(v, c["Cliente"])]
            if achados:
                break
    achados = list(dict.fromkeys(achados))
    if not achados:
        return False

    def status(a):
        return classe_sgp(sgp.contratos[a[1]]["Status"]) if a[0] == "SGP" else \
            classe_status(ixc.contrato(a[1])["Status acesso"])

    def nome(a):
        c = sgp.contratos[a[1]] if a[0] == "SGP" else ixc.contrato(a[1])
        return f"{a[0]} {a[1]} {c['Cliente']} ({c['Status'] if a[0] == 'SGP' else c['Status acesso']})"

    # vários contratos do mesmo cliente (ex.: um antigo desativado e um ativo): vale o ativo
    if len({compacto(nome(a).split(" ", 2)[2].rsplit(" (", 1)[0]) for a in achados}) == 1:
        ativos = [a for a in achados if status(a) == "Ativo"]
        if ativos:
            achados = ativos
    if len({status(a) for a in achados}) > 1:
        o["classe"] = "Conferir"
        o["obs"].append(f"{por} bate com mais de um cliente: " + "; ".join(nome(a) for a in achados[:5]))
        return True
    sistema, cid = achados[0]
    if len(achados) > 1:
        o["obs"].append(f"{por} bate com {len(achados)} contratos, todos {status(achados[0]).lower()}: "
                        + "; ".join(nome(a) for a in achados[:5]))
    if sistema == "SGP":
        o["sgp_pendente"] = cid
        o["achado_por"] = f"{por} (SGP)"
        o["classe"] = status(achados[0])
        parceiro = ixc.logins.get(PARCEIRO_SGP)
        if parceiro:
            c = ixc.contrato(parceiro[0]["ID contrato"])
            o.update(contrato=parceiro[0]["ID contrato"], cliente=c["Cliente"], plano=c["Descrição"])
    else:
        c = ixc.contrato(cid)
        o.update(contrato=cid, achado_por=f"{por} (IXC)", cliente=c["Cliente"], plano=c["Descrição"],
                 plano_id=c["Plano de venda"], status_ixc=c["Status acesso"], classe=status(achados[0]))
    return True


def classe_sgp(status):
    s = chave(status)
    if s in SGP_ATIVO:
        return "Ativo"
    if s in SGP_DESATIVADO:
        return "Desativado"
    if s in SGP_PRE:
        return "Pré-contrato"
    return "Conferir"


def fmt_sgp(cs):
    return "; ".join(f"{c['Contrato ID']} {c['Cliente']} ({c['Status']})" for c in cs[:5])


def classificar_sgp(o, sgp):
    """Substitui o status do contrato do parceiro pelo status do cliente dele no SGP."""
    o["tipo"] = f"Rede neutra ({NOME_PARCEIRO_SGP})"
    m = re.match(r"^(\d+)\s*[-_ ]*\s*(.*)$", o["sufixo"])
    cid, nome = (int(m.group(1)), m.group(2)) if m else (None, o["sufixo"])
    if cid is not None:
        c = sgp.contratos.get(cid)
        if not c:
            o["classe"] = "Desativado" if o["estado"] == "Down" else "Não encontrado"
            o["obs"].append(f"Contrato {cid} não existe no SGP {NOME_PARCEIRO_SGP}")
            return
        candidatos = [c]
        o["sgp_achado_por"] = "Contrato SGP"
    else:
        candidatos = sgp.por_nome(nome)
        o["sgp_achado_por"] = "Nome"
        if not candidatos:
            o["classe"] = "Desativado" if o["estado"] == "Down" else "Conferir"
            o["obs"].append(f"Sem número do contrato SGP e nenhum cliente {NOME_PARCEIRO_SGP} com o nome '{nome}'")
            return
        classes = {classe_sgp(c["Status"]) for c in candidatos}
        if len(candidatos) > 1 and len(classes) > 1:
            o["classe"] = "Conferir"
            o["obs"].append(f"Nome '{nome}' bate com vários clientes {NOME_PARCEIRO_SGP}: {fmt_sgp(candidatos)}")
            return
    c = candidatos[0]
    o.update(sgp_contrato=", ".join(str(x["Contrato ID"]) for x in candidatos),
             sgp_cliente="; ".join(dict.fromkeys(x["Cliente"] for x in candidatos)),
             sgp_status=c["Status"], status_ixc=f"{c['Status']} (SGP)", classe=classe_sgp(c["Status"]))
    if len(candidatos) > 1:
        o["obs"].append(f"Nome '{nome}' bate com {len(candidatos)} contratos, todos {o['classe'].lower()}")
    if o["classe"] == "Desativado":
        outros = [x for x in sgp.por_cliente[c["Cliente ID"]]
                  if x not in candidatos and classe_sgp(x["Status"]) == "Ativo"]
        if outros:
            o["classe"] = "Conferir"
            o["obs"].append(f"Contrato cancelado, mas o cliente tem outro contrato ativo no SGP: {fmt_sgp(outros)}")


# ---------------------------------------------------------------- classificação

def extrair_numero(descricao):
    """CTR-555, CTR-555-NOME, C-TR-555, CTR555, CTR.555.NOME, UPWIRE-CTR-555, 555-NOME → '555'."""
    m = re.match(r"^(?:[A-Z]+-)?\s*C-?(?:TR|RT)\s*[-_. ]?\s*(\d+)(.*)$", descricao.strip(), re.I) or \
        re.match(r"^\s*(\d+)(.*)$", descricao)
    if not m:
        return None, descricao
    return m.group(1), m.group(2).strip(" -_.")


def classe_status(status_acesso):
    s = chave(status_acesso)
    if s in STATUS_ATIVO:
        return "Ativo"
    if s in STATUS_DESATIVADO:
        return "Desativado"
    if s in STATUS_PRE:
        return "Pré-contrato"
    return "Conferir"


def compacto(txt):
    txt = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", txt.upper())


def sugerir_cliente(ixc, descricao):
    """Para descrições sem número: procura clientes cujo nome contenha uma palavra da descrição."""
    partes = [compacto(p) for p in re.split(r"[-_. ]+", descricao) if p.upper() not in ("CTR", "C")]
    partes = [p for p in partes if len(p) >= 5] or [compacto(descricao.replace("CTR", ""))]
    achados = {}
    for cid, c in ixc.contratos.items():
        nome = compacto(c["Cliente"])
        if any(p and len(p) >= 5 and (p in nome or (len(nome) >= 8 and nome in p)) for p in partes):
            achados.setdefault(c["Cliente"], []).append((cid, c["Status acesso"]))
    return "; ".join(f"{nome} (contrato {', '.join(f'{i}/{s}' for i, s in v[:3])})"
                     for nome, v in list(achados.items())[:3])


def classificar(onus, ixc, sgp=None, outras_olts=None):
    repeticoes = collections.Counter(extrair_numero(o["descricao"])[0] for o in onus)
    descricoes = collections.Counter(o["descricao"].upper().replace("ID", "") for o in onus)

    for o in onus:
        num, sufixo = extrair_numero(o["descricao"])
        o.update(numero=num, sufixo=sufixo, contrato=None, login=None, cliente=None, plano=None,
                 status_ixc=None, login_ativo=None, achado_por=None, obs=[])

        if num is None:
            desc = compacto(o["descricao"])
            cid = next((i for nome, i in NOMES_CONHECIDOS if nome in desc), None)
            c = ixc.contrato(cid)
            if c:
                o.update(contrato=cid, achado_por="Nome", cliente=c["Cliente"], plano=c["Descrição"],
                         plano_id=c["Plano de venda"], status_ixc=c["Status acesso"])
                o["classe"] = classe_status(c["Status acesso"])
            elif identificar_por_nome(o, ixc, sgp):
                pass
            elif o["estado"] == "Down":
                o["classe"] = "Desativado"
                o["obs"].append("Descrição sem número e sem cliente identificado, ONU offline")
            else:
                o["classe"] = "Conferir"
                o["obs"].append("Descrição sem número e sem cliente identificado, ONU " +
                                ("online" if o["estado"] == "Up" else "com estado desconhecido"))
                if "NOVASOLU" in desc:
                    o["obs"].append("Ex-Nova Solutec: pode ser cliente CCS ou Upwire e estar configurado por VLAN "
                                    "(/30 ou /29), sem login PPPoE")
                sug = sugerir_cliente(ixc, o["descricao"])
                if sug:
                    o["obs"].append(f"Possível cliente: {sug}")
        else:
            por_login = ixc.logins.get(num, [])
            por_contrato = ixc.contrato(num)
            cids_login = {r["ID contrato"] for r in por_login}

            if por_login and por_contrato and int(num) not in cids_login:
                c_login = ixc.contrato(next(iter(cids_login)))
                mesmo_cliente = c_login and compacto(c_login["Cliente"]) == compacto(por_contrato["Cliente"])
                if not mesmo_cliente:
                    o["obs"].append(f"Número {num} é login do contrato {sorted(cids_login)} e também é o "
                                    f"contrato {num} ({por_contrato['Cliente']}) — conferir")
                    o["conflito"] = True

            if por_login:
                r = por_login[0]
                c = ixc.contrato(r["ID contrato"])
                o.update(contrato=r["ID contrato"], login=r["Login"], achado_por="Login",
                         login_ativo=r.get("Ativo"), ult_conexao=r.get("Última conexão inicial"))
                if len(cids_login) > 1:
                    o["obs"].append(f"Login existe em mais de um contrato: {sorted(cids_login)}")
            elif por_contrato:
                c = por_contrato
                o.update(contrato=num and int(num), achado_por="Contrato")
            elif num in NUMEROS_ANTIGOS:
                c = ixc.contrato(NUMEROS_ANTIGOS[num])
                o.update(contrato=NUMEROS_ANTIGOS[num], achado_por="Nº antigo")
            else:
                c = None

            if c:
                o.update(cliente=c["Cliente"], plano=c["Descrição"], plano_id=c["Plano de venda"],
                         status_ixc=c["Status acesso"], status_contrato=c["Status contrato"])
                o["classe"] = "Conferir" if o.get("conflito") else classe_status(c["Status acesso"])
            else:
                o["obs"].append("Número não existe no IXC como login nem como contrato"
                                " (nº do sistema antigo da CCS Brasil)")
                o["classe"] = "Desativado" if o["estado"] == "Down" else "Não encontrado"

        # tipo de cliente
        rep = repeticoes[num] if num else 0
        if o.get("plano_id") == PLANO_REDE_NEUTRA or o.get("contrato") in PARCEIROS_EXTRA:
            o["tipo"] = "Rede neutra"
        elif rep >= 3:
            o["tipo"] = "Rede neutra (provável)"
            o["obs"].append(f"Mesmo número em {rep} ONUs desta OLT")
        else:
            o["tipo"] = "Cliente final"
            if rep == 2:
                o["obs"].append("Mesmo número em 2 ONUs desta OLT")

        if sgp and num == PARCEIRO_SGP:
            classificar_sgp(o, sgp)
        elif o.get("sgp_pendente"):
            o["sufixo"] = str(o["sgp_pendente"])
            classificar_sgp(o, sgp)
        elif sgp and num and len(num) == 5 and o.get("achado_por") is None and int(num) in sgp.contratos:
            # cliente Upwire sem o 201939 na descrição (ex.: CTR-11267)
            o["obs"] = [f"Descrição sem o {PARCEIRO_SGP}: contrato {num} encontrado no SGP {NOME_PARCEIRO_SGP}"]
            o.update(sufixo=num + ("-" + o["sufixo"] if o["sufixo"] else ""))
            classificar_sgp(o, sgp)

        if o["descricao"] and descricoes[o["descricao"].upper().replace("ID", "")] > 1:
            o["obs"].append("Descrição duplicada nesta OLT (possível ONU antiga não removida)")

    # ONU antiga trocada: mesmo login (cliente final) ou mesma descrição (rede neutra) online em outra ONU
    def grupo(o):
        if o.get("sgp_contrato") and "," not in str(o["sgp_contrato"]):
            return ("sgp", str(o["sgp_contrato"]))
        if o["tipo"] == "Cliente final" and o.get("login") is not None:
            return ("login", chave(o["login"]))
        return ("desc", o["descricao"].upper().replace("ID", "").rstrip(" -_."))
    online = collections.defaultdict(list)
    for o in onus:
        if o["estado"] == "Up" and o["descricao"]:
            online[grupo(o)].append(o)
    for o in onus:
        outras = online.get(grupo(o))
        if o["estado"] == "Down" and o["descricao"] and outras and o["classe"] in ("Ativo", "Conferir", "Não encontrado"):
            o["classe"] = "Desativado"
            o["obs"].append("ONU antiga: o mesmo cliente está online em " +
                            ", ".join(f"{x['pon']} ONU {x['onu_id']} ({x['descricao']})" for x in outras))

    # mesmo serial provisionado em outra OLT: se aqui está offline e lá online, a ONU mudou de lugar
    for o in onus:
        for olt, x in (outras_olts or {}).get(o["serial"], []):
            o["obs"].append(f"Serial também provisionado na {olt} ({x['descricao']}, {x['estado']})")
            if o["estado"] == "Down" and x["estado"] == "Up" and o["classe"] != "Desativado":
                o["classe"] = "Desativado"
                o["obs"].append("ONU antiga: este equipamento está online em outra OLT")

    for o in onus:
        if (o["serial"], o["descricao"]) in DECISOES_MANUAIS:
            o["classe"], motivo = DECISOES_MANUAIS[(o["serial"], o["descricao"])]
            o["obs"] = [motivo]

    for o in onus:
        if o.get("contrato") is None and any(n in compacto(o["descricao"]) for n in REMOVER_SEM_CONTRATO):
            o["classe"] = "Desativado"
            o["obs"] = [f"{o['descricao']}: remover (confirmado pelo usuário; sem contrato de rede neutra)"]

    for o in onus:  # ambíguo ou não identificado: offline = remover; online continua em Conferir
        if o["classe"] in ("Conferir", "Não encontrado") and o["estado"] == "Down":
            o["classe"] = "Desativado"
            o["obs"].append("Não foi possível confirmar o cliente e a ONU está offline")

    for o in onus:  # alertas de estado da ONU
        if o.get("perfil"):
            o["obs"].append(f"Perfil: {o['perfil']}")
        if o.get("sinal") is not None and o["sinal"] <= -27:
            o["obs"].append(f"Sinal óptico baixo: {o['sinal']} dBm")
        if o["classe"] == "Desativado" and o["estado"] == "Up":
            o["obs"].append("ONU ONLINE com contrato desativado — confirmar antes de remover")
        if o["classe"] == "Ativo" and o["estado"] == "Down":
            ult = o.get("ult_conexao")
            o["obs"].append("Cliente ativo com ONU offline" +
                            (f" (última conexão PPPoE: {ult:%d/%m/%Y})" if hasattr(ult, "strftime") else ""))
        if o["classe"] == "Ativo" and o.get("login_ativo") == "Não":
            o["obs"].append("Contrato ativo, mas o login PPPoE está desabilitado no IXC")
        if o["classe"] == "Ativo" and not chave(o.get("status_ixc")).startswith("ativo"):
            o["obs"].append(f"Status: {o['status_ixc']}")
    return onus


# ---------------------------------------------------------------- Excel

AZUL = PatternFill("solid", start_color="1F4E78")
CINZA = PatternFill("solid", start_color="F2F2F2")
AMARELO = PatternFill("solid", start_color="FFFF00")
CORES = {"Ativo": "E2EFDA", "Desativado": "FCE4D6", "Conferir": "FFF2CC",
         "Não encontrado": "FFF2CC", "Pré-contrato": "DDEBF7"}

COLUNAS = [
    ("PON", "pon", 9), ("ONU ID", "onu_id", 8), ("Serial", "serial", 15), ("Estado ONU", "estado", 11),
    ("Descrição na OLT", "descricao", 38), ("Contrato IXC", "contrato", 11), ("Login PPPoE", "login", 16),
    ("Cliente", "cliente", 40), ("Plano", "plano", 40), ("Tipo", "tipo", 20),
    (f"Contrato {NOME_PARCEIRO_SGP} (SGP)", "sgp_contrato", 14), (f"Cliente {NOME_PARCEIRO_SGP}", "sgp_cliente", 36),
    ("Status acesso IXC", "status_ixc", 20), ("Classificação", "classe", 15),
    ("Achado por", "achado_por", 11), ("Observações", "obs", 70),
]


def letra(campo, colunas=COLUNAS):
    return get_column_letter([k for _, k, _ in colunas].index(campo) + 1)


def fmt_cab(ws, linha, ncols):
    for c in range(1, ncols + 1):
        cel = ws.cell(linha, c)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.fill = AZUL
        cel.alignment = Alignment(vertical="center", wrap_text=True)


def aba_lista(wb, titulo, onus, colunas=COLUNAS):
    ws = wb.create_sheet(titulo)
    for i, (nome, _, larg) in enumerate(colunas, 1):
        ws.cell(1, i, nome)
        ws.column_dimensions[get_column_letter(i)].width = larg
    fmt_cab(ws, 1, len(colunas))
    col_classe = [k for _, k, _ in colunas].index("classe") + 1
    for r, o in enumerate(onus, 2):
        for c, (_, campo, _) in enumerate(colunas, 1):
            v = o.get(campo)
            if isinstance(v, list):
                v = " | ".join(v)
            ws.cell(r, c, v).font = Font(name=FONTE)
        if o.get("estado") == "Down":
            ws.cell(r, [k for _, k, _ in colunas].index("estado") + 1).font = Font(name=FONTE, bold=True, color="C00000")
        cor = CORES.get(o.get("classe"))
        if cor:
            ws.cell(r, col_classe).fill = PatternFill("solid", start_color=cor)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(len(onus) + 1, 1)}"
    return ws


def titulo(ws, texto, sub):
    ws["A1"] = texto
    ws["A1"].font = Font(name=FONTE, bold=True, size=14)
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONTE, italic=True, color="595959")


def ordenar(onus):
    return sorted(onus, key=lambda o: (tuple(int(x) for x in o["pon"].split("/")), o["onu_id"]))


def aba_pop(wb, nome_olt, onus, pop):
    """Inventário do POP (dados/pops.json) e ocupação das PONs."""
    ws = wb.create_sheet("POP")
    titulo(ws, f"{nome_olt} — Equipamentos do POP", pop.get("pop", ""))
    cab = ["Equipamento", "Fabricante", "Modelo", "Observação"]
    for c, v in enumerate(cab, 1):
        ws.cell(4, c, v)
    fmt_cab(ws, 4, len(cab))
    r = 5
    for e in pop.get("equipamentos", []):
        for c, k in enumerate(("equipamento", "fabricante", "modelo", "obs"), 1):
            cel = ws.cell(r, c, e.get(k))
            cel.font = Font(name=FONTE)
            if e.get(k) is None or (k == "modelo" and "definir" in str(e.get(k))):
                cel.fill = AMARELO
        r += 1
    r += 2
    ws.cell(r - 1, 1, "Ocupação das PONs (todas as ONUs provisionadas)").font = Font(name=FONTE, bold=True)
    cab = ["PON", "ONUs provisionadas", "Online (Up)", "Offline (Down)", "A remover"]
    for c, v in enumerate(cab, 1):
        ws.cell(r, c, v)
    fmt_cab(ws, r, len(cab))
    pons = sorted({o["pon"] for o in onus}, key=lambda p: tuple(int(x) for x in p.split("/")))
    for i, pon in enumerate(pons, r + 1):
        do_pon = [o for o in onus if o["pon"] == pon]
        valores = [pon, len(do_pon), sum(o["estado"] == "Up" for o in do_pon),
                   sum(o["estado"] == "Down" for o in do_pon), sum(o["classe"] == "Desativado" for o in do_pon)]
        for c, v in enumerate(valores, 1):
            ws.cell(i, c, v).font = Font(name=FONTE)
    total_pon = pop.get("portas_pon")
    if total_pon:
        ws.cell(r + len(pons) + 2, 1, f"PONs em uso: {len(pons)} de {total_pon}").font = Font(name=FONTE, bold=True)
    n = pop.get("colocation_linhas", 0)
    if n:
        r = r + len(pons) + 5
        ws.cell(r - 1, 1, "Colocation — equipamentos de clientes no rack (preencher)").font = Font(name=FONTE, bold=True)
        cab = ["Cliente", "Equipamento", "Modelo", "Posição no rack (U)", "Observação"]
        for c, v in enumerate(cab, 1):
            ws.cell(r, c, v)
        fmt_cab(ws, r, len(cab))
        for i in range(r + 1, r + 1 + n):
            for c in range(1, len(cab) + 1):
                ws.cell(i, c).fill = AMARELO
        ws.cell(r + n + 2, 1, "Células amarelas: preencher manualmente.").font = Font(name=FONTE, italic=True, color="595959")
    for c, w in zip("ABCDE", (22, 22, 34, 40, 30)):
        ws.column_dimensions[c].width = w


def gerar_limpeza(nome_olt, onus, destino):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    titulo(ws, f"{nome_olt} — Limpeza de ONUs", "Cruzamento da descrição das ONUs com os relatórios do IXC")
    linhas = [("Classificação", "Qtd. ONUs", "O que fazer"),
              ("Desativado", None, "Remover da OLT (aba 'Remover'). Veja a coluna Observações antes."),
              ("Não encontrado", None, "Número não existe no IXC. Verificar e, se for o caso, remover."),
              ("Conferir", None, "Descrição sem número ou ambígua. Identificar o cliente."),
              ("Pré-contrato", None, "Aguardando assinatura. Manter."),
              ("Ativo", None, "Manter. Vão para o relatório de ativos."),
              ("Total", None, None)]
    for i, (a, b, c) in enumerate(linhas, 4):
        ws.cell(i, 1, a); ws.cell(i, 3, c)
        if i == 4:
            ws.cell(i, 2, b)
            fmt_cab(ws, 4, 3)
        elif a == "Total":
            ws.cell(i, 2, f"=SUM(B5:B{i - 1})")
            for c_ in (1, 2):
                ws.cell(i, c_).font = Font(name=FONTE, bold=True)
        else:
            ws.cell(i, 2, f"=COUNTIF(Todas!{letra('classe')}:{letra('classe')},A{i})")
            ws.cell(i, 1).fill = PatternFill("solid", start_color=CORES[a])
        for c_ in (1, 2, 3):
            if i > 4 and a != "Total":
                ws.cell(i, c_).font = Font(name=FONTE)
    ws.cell(12, 1, "Regras usadas").font = Font(name=FONTE, bold=True)
    regras = [
        "O número após 'CTR-' é procurado primeiro como login PPPoE e depois como ID de contrato.",
        "Ativo = status de acesso Ativo, Bloqueio Automático/Manual ou Financeiro em atraso (bloqueado conta como ativo).",
        f"Rede neutra = plano de venda {PLANO_REDE_NEUTRA} (Circuito Dedicado). 'Provável' = mesmo número em 3 ou mais ONUs.",
        "Uma ONU de rede neutra segue o status do contrato do provedor parceiro.",
        "Número que não existe no IXC (sistema antigo) ou descrição sem cliente identificado: offline = remover; online = conferir.",
        "ONU offline cujo mesmo login (ou mesma descrição, na rede neutra) está online em outra ONU da OLT = ONU antiga, remover.",
        f"Exceção {NOME_PARCEIRO_SGP} (CTR-{PARCEIRO_SGP}-...): segue o status do cliente no SGP do parceiro, "
        "buscado pelo nº do contrato SGP e, sem número, pelo nome. Cancelado = remover; Suspenso = ativo; "
        "Inativo (nunca ativado) = pré-contrato.",
        f"Fonte: relatórios de contrato e de login do IXC (CCS Brasil) e de contratos do SGP {NOME_PARCEIRO_SGP} (fornecidos pelo usuário).",
    ]
    for i, r in enumerate(regras, 13):
        ws.cell(i, 1, f"• {r}").font = Font(name=FONTE)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 70

    remover = [o for o in onus if o["classe"] == "Desativado"]
    conferir = [o for o in onus if o["classe"] in ("Conferir", "Não encontrado")]
    aba_lista(wb, "Remover", remover)
    aba_lista(wb, "Conferir", conferir)
    aba_lista(wb, "Todas", onus)
    wb.save(destino)


def gerar_ativos(nome_olt, onus, destino, pop=None):
    ativos = [o for o in onus if o["classe"] in ("Ativo", "Pré-contrato")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    titulo(ws, f"{nome_olt} — Clientes ativos",
           "ONUs com contrato ativo no IXC (bloqueados contam como ativos). Desativados e não identificados ficam fora.")

    colunas = [c for c in COLUNAS if c[1] not in ("achado_por",)]
    T, E, K = letra("tipo", colunas), letra("estado", colunas), letra("contrato", colunas)
    pons = sorted({o["pon"] for o in ativos}, key=lambda p: tuple(int(x) for x in p.split("/")))
    cab = ["PON", "ONUs ativas", "Cliente final", "Rede neutra", "Online (Up)", "Offline (Down)"]
    for c, v in enumerate(cab, 1):
        ws.cell(4, c, v)
    fmt_cab(ws, 4, len(cab))
    for i, pon in enumerate(pons, 5):
        ws.cell(i, 1, pon)
        ws.cell(i, 2, f"=COUNTIF(Ativos!A:A,A{i})")
        ws.cell(i, 3, f'=COUNTIFS(Ativos!A:A,A{i},Ativos!{T}:{T},"Cliente final")')
        ws.cell(i, 4, f'=COUNTIFS(Ativos!A:A,A{i},Ativos!{T}:{T},"Rede neutra*")')
        ws.cell(i, 5, f'=COUNTIFS(Ativos!A:A,A{i},Ativos!{E}:{E},"Up")')
        ws.cell(i, 6, f'=COUNTIFS(Ativos!A:A,A{i},Ativos!{E}:{E},"Down")')
    tot = 5 + len(pons)
    ws.cell(tot, 1, "Total")
    for c in range(2, 7):
        ws.cell(tot, c, f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}{tot - 1})")
    for r in range(5, tot + 1):
        for c in range(1, 7):
            ws.cell(r, c).font = Font(name=FONTE, bold=(r == tot))
    for c, w in zip("ABCDEF", (12, 13, 13, 13, 13, 14)):
        ws.column_dimensions[c].width = w

    # ranking por contrato/provedor
    r0 = tot + 3
    ws.cell(r0 - 1, 1, "ONUs por cliente / provedor").font = Font(name=FONTE, bold=True)
    for c, v in enumerate(["Contrato", "Cliente / provedor", "Tipo", "ONUs"], 1):
        ws.cell(r0, c, v)
    fmt_cab(ws, r0, 4)
    grupos = collections.Counter((o["contrato"], o["cliente"], o["tipo"]) for o in ativos)
    for i, ((cid, cli, tipo), _) in enumerate(sorted(grupos.items(), key=lambda kv: -kv[1]), r0 + 1):
        ws.cell(i, 1, cid); ws.cell(i, 2, cli); ws.cell(i, 3, tipo)
        ws.cell(i, 4, f"=COUNTIF(Ativos!{K}:{K},A{i})")
        for c in range(1, 5):
            ws.cell(i, c).font = Font(name=FONTE)
    ws.column_dimensions["B"].width = 45

    aba_lista(wb, "Ativos", ordenar(ativos), colunas)
    if pop:
        aba_pop(wb, nome_olt, onus, pop)
    wb.save(destino)
    return ativos


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("olt", type=Path)
    ap.add_argument("--logins", type=Path, required=True)
    ap.add_argument("--contratos", type=Path, required=True)
    ap.add_argument("--saida", type=Path, default=Path("saida"))
    ap.add_argument("--sgp", type=Path, help=f"Relatório de contratos do SGP da {NOME_PARCEIRO_SGP}")
    ap.add_argument("--fabricante", choices=sorted(LEITORES))
    ap.add_argument("--nome", help="Nome da OLT (padrão: nome do arquivo)")
    ap.add_argument("--pops", type=Path, default=Path("dados/pops.json"),
                    help="Inventário dos POPs: {nome_olt: {pop, portas_pon, equipamentos: [...]}}")
    a = ap.parse_args()

    texto = a.olt.read_text(encoding="utf-8", errors="replace")
    onus = LEITORES[a.fabricante or detectar_fabricante(texto)](texto)
    if not onus:
        raise SystemExit("Nenhuma ONU encontrada no arquivo.")
    nome = a.nome or a.olt.stem

    # outras OLTs na mesma pasta, para achar o mesmo serial provisionado em mais de um lugar
    outras = collections.defaultdict(list)
    for arq in sorted(a.olt.parent.glob("*.txt")):
        if arq.resolve() != a.olt.resolve():
            t = arq.read_text(encoding="utf-8", errors="replace")
            try:
                for x in LEITORES[detectar_fabricante(t)](t):
                    outras[x["serial"]].append((arq.stem, x))
            except SystemExit:
                pass

    onus = ordenar(classificar(onus, IXC(a.logins, a.contratos), SGP(a.sgp) if a.sgp else None, outras))
    pop = json.loads(a.pops.read_text(encoding="utf-8")).get(nome) if a.pops.exists() else None

    a.saida.mkdir(parents=True, exist_ok=True)
    gerar_limpeza(nome, onus, a.saida / f"{nome}_limpeza.xlsx")
    gerar_ativos(nome, onus, a.saida / f"{nome}_ativos.xlsx", pop)
    cont = collections.Counter(o["classe"] for o in onus)
    print(f"{nome}: {len(onus)} ONUs — " + ", ".join(f"{k}: {v}" for k, v in cont.most_common()))


if __name__ == "__main__":
    main()
